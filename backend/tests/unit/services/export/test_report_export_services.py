"""Testy jednostkowe renderów eksportu raportów — Story 6.7 (AC1-AC6)."""

import io
import json

import pytest
import yaml
from openpyxl import load_workbook

from app.services.export import (
    csv_export_service,
    excel_export_service,
    html_export_service,
    json_export_service,
    pdf_export_service,
    yaml_export_service,
)


@pytest.fixture()
def report() -> dict:
    return {
        "project_id": "11111111-2222-3333-4444-555555555555",
        "generated_at": "2026-07-23T10:00:00Z",
        "kpis": {
            "compliance_score": 72.4,
            "coverage_percentage": 61.0,
            "total_findings": 2,
            "critical_count": 1,
            "high_count": 0,
            "medium_count": 1,
            "low_count": 0,
            "test_cases_total": 5,
            "test_cases_ready": 2,
            "test_cases_passed": 1,
        },
        "severity_counts": {"critical": 1, "high": 0, "medium": 1, "low": 0},
        "module_scores": [
            {"file_path": "app/db.py", "score": 40.0, "total_findings": 1},
        ],
        "trend_points": [
            {"label": "Analiza #1", "score": 50.0, "timestamp": None},
            {"label": "Analiza #2", "score": 72.0, "timestamp": None},
        ],
        "conclusions": [
            "Zgodność oprogramowania z dokumentacją wynosi 72% — poziom wymagający uwagi.",
            "Wykryto 1 problemów o znaczeniu krytycznym — zalecamy ich naprawę w pierwszej kolejności.",
        ],
        "findings": [
            {
                "id": "sf1",
                "source": "security",
                "title": "SQL Injection w zapytaniu użytkownika",
                "description": "Zapytanie budowane przez konkatenację — znaki < > & \" muszą przeżyć eksport.",
                "severity": "critical",
                "cwe_id": "CWE-89",
                "cvss_score": 9.8,
                "file_path": "app/db.py",
                "line_start": 10,
                "code_snippet": 'cursor.execute("SELECT * FROM users WHERE id = " + user_id)',
                "language": "python",
                "recommendation": "Użyć zapytań parametryzowanych; nie sklejać SQL ze stringów.",
            },
            {
                "id": "gf1",
                "source": "gap",
                "title": "Brak walidacji formularza",
                "description": "Wymaganie: system waliduje dane wejściowe (ąęłśżź).",
                "severity": "medium",
                "cwe_id": None,
                "cvss_score": None,
                "file_path": "app/forms.py",
                "line_start": None,
                "code_snippet": None,
                "language": None,
                "recommendation": None,
            },
        ],
    }


class TestCsvExport:
    def test_bom_semicolon_and_polish_headers(self, report):
        data = csv_export_service.render(report)
        text = data.decode("utf-8")
        assert text.startswith("\ufeff")
        header = text.splitlines()[0].lstrip("\ufeff")
        assert header.split(";")[:3] == ["Lp.", "Źródło", "Tytuł"]
        assert "Istotność" in header
        assert "Krytyczny" in text
        assert "Skan bezpieczeństwa" in text

    def test_escapes_semicolons_in_content(self, report):
        report["findings"][0]["title"] = "Tytuł; ze średnikiem"
        text = csv_export_service.render(report).decode("utf-8")
        assert '"Tytuł; ze średnikiem"' in text


class TestJsonExport:
    def test_pretty_snake_case_utf8(self, report):
        data = json_export_service.render(report, "executive")
        text = data.decode("utf-8")
        parsed = json.loads(text)
        assert parsed["report_type"] == "executive"
        assert parsed["kpis"]["compliance_score"] == 72.4
        assert parsed["findings"][0]["cwe_id"] == "CWE-89"
        # pretty-printed + polskie znaki bez escapowania
        assert "\n  " in text
        assert "ąęłśżź" in text


class TestYamlExport:
    def test_round_trip_and_unicode(self, report):
        data = yaml_export_service.render(report, "both")
        text = data.decode("utf-8")
        assert "ąęłśżź" in text
        parsed = yaml.safe_load(text)
        assert parsed["report_type"] == "both"
        assert parsed["kpis"]["total_findings"] == 2
        assert parsed["findings"][1]["severity"] == "medium"


class TestExcelExport:
    def test_four_sheets_with_polish_names(self, report):
        data = excel_export_service.render(report)
        wb = load_workbook(io.BytesIO(data))
        assert wb.sheetnames == ["Podsumowanie", "Ustalenia", "Przypadki testowe", "Pokrycie"]

    def test_findings_sheet_content_and_severity_fill(self, report):
        wb = load_workbook(io.BytesIO(excel_export_service.render(report)))
        ws = wb["Ustalenia"]
        assert ws.cell(row=1, column=5).value == "Istotność"
        assert ws.cell(row=1, column=1).font.bold is True
        assert ws.cell(row=2, column=5).value == "Krytyczny"
        # Wzorek (nie tylko kolor) — WCAG
        assert ws.cell(row=2, column=5).fill.fill_type == "darkDown"

    def test_summary_and_coverage_sheets(self, report):
        wb = load_workbook(io.BytesIO(excel_export_service.render(report)))
        summary_values = [row[0].value for row in wb["Podsumowanie"].iter_rows()]
        assert "Zgodność z dokumentacją" in summary_values
        coverage = wb["Pokrycie"]
        assert coverage.cell(row=2, column=1).value == "app/db.py"


class TestHtmlExport:
    def test_standalone_html_with_inline_svg(self, report):
        text = html_export_service.render(report, "both").decode("utf-8")
        assert text.startswith("<!DOCTYPE html>")
        assert "<style>" in text
        assert "<svg" in text and "<pattern" in text
        assert "Kluczowe wnioski" in text
        assert "CWE-89" in text
        # Autoescape — HTML z opisu nie może przeciec jako znaczniki
        assert "&lt;" in text or "<" not in report["findings"][0]["description"]

    def test_executive_type_hides_findings(self, report):
        text = html_export_service.render(report, "executive").decode("utf-8")
        assert "Szczegółowe ustalenia" not in text
        assert "Podsumowanie wykonawcze" in text


class TestPdfExport:
    def test_generates_valid_pdf(self, report):
        data = pdf_export_service.render(report, "both")
        assert data.startswith(b"%PDF")
        assert len(data) > 1000

    def test_executive_pdf_smaller_than_full(self, report):
        # Duplikacja ustaleń, by różnica była mierzalna
        report["findings"] = report["findings"] * 20
        full = pdf_export_service.render(report, "both")
        executive = pdf_export_service.render(report, "executive")
        assert len(executive) < len(full)
