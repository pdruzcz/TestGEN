"""Eksport Excel (.xlsx) — Story 6.7, AC4.

openpyxl, cztery arkusze: Podsumowanie, Ustalenia, Przypadki testowe,
Pokrycie. Nagłówki bold, auto-szerokość kolumn, wyróżnienie severity
wzorkiem I kolorem (WCAG — nie tylko kolor).
"""

import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from app.services.export.common import kpi_rows, severity_pl, source_pl

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill(start_color="DDE7F5", end_color="DDE7F5", fill_type="solid")

# Wzorek + kolor per severity (openpyxl pattern fills) — AC6/WCAG.
SEVERITY_FILLS: dict[str, PatternFill] = {
    "critical": PatternFill(start_color="C62828", end_color="FFFFFF", fill_type="darkDown"),
    "high": PatternFill(start_color="EF6C00", end_color="FFFFFF", fill_type="lightUp"),
    "medium": PatternFill(start_color="F9A825", end_color="FFFFFF", fill_type="lightGrid"),
    "low": PatternFill(start_color="2E7D32", end_color="FFFFFF", fill_type="lightHorizontal"),
}

MAX_COL_WIDTH = 60


def _autosize(ws: Worksheet) -> None:
    for column_cells in ws.columns:
        longest = max((len(str(c.value)) for c in column_cells if c.value is not None), default=0)
        letter = column_cells[0].column_letter
        ws.column_dimensions[letter].width = min(MAX_COL_WIDTH, max(10, longest + 2))


def _write_header(ws: Worksheet, headers: list[str]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center")


def _sheet_summary(ws: Worksheet, report: dict) -> None:
    _write_header(ws, ["Wskaźnik", "Wartość"])
    for label, value in kpi_rows(report):
        ws.append([label, value])
    ws.append([])
    ws.append(["Kluczowe wnioski", ""])
    ws.cell(row=ws.max_row, column=1).font = HEADER_FONT
    for conclusion in report.get("conclusions") or []:
        ws.append([conclusion, ""])
    ws.append([])
    ws.append(["Data wygenerowania", report.get("generated_at") or ""])
    _autosize(ws)


def _sheet_findings(ws: Worksheet, report: dict) -> None:
    _write_header(
        ws,
        ["Lp.", "Źródło", "Tytuł", "Opis", "Istotność", "CWE", "CVSS", "Plik", "Linia", "Rekomendacja"],
    )
    for i, f in enumerate(report.get("findings") or [], start=1):
        severity = (f.get("severity") or "").lower()
        ws.append(
            [
                i,
                source_pl(f.get("source")),
                f.get("title") or "",
                f.get("description") or "",
                severity_pl(severity),
                f.get("cwe_id") or "",
                f.get("cvss_score"),
                f.get("file_path") or "",
                f.get("line_start"),
                f.get("recommendation") or "",
            ]
        )
        fill = SEVERITY_FILLS.get(severity)
        if fill is not None:
            ws.cell(row=ws.max_row, column=5).fill = fill
    _autosize(ws)


def _sheet_test_cases(ws: Worksheet, report: dict) -> None:
    kpis = report.get("kpis") or {}
    _write_header(ws, ["Wskaźnik", "Wartość"])
    ws.append(["Przypadki testowe (łącznie)", kpis.get("test_cases_total", 0)])
    ws.append(["Gotowe do uruchomienia", kpis.get("test_cases_ready", 0)])
    ws.append(["Zakończone pozytywnie", kpis.get("test_cases_passed", 0)])
    _autosize(ws)


def _sheet_coverage(ws: Worksheet, report: dict) -> None:
    kpis = report.get("kpis") or {}
    _write_header(ws, ["Moduł (plik)", "Zgodność [%]", "Liczba problemów"])
    for module in report.get("module_scores") or []:
        ws.append(
            [
                module.get("file_path") or "",
                round(module.get("score") or 0),
                module.get("total_findings", 0),
            ]
        )
    ws.append([])
    coverage = kpis.get("coverage_percentage")
    ws.append(["Pokrycie wymagań biznesowych", f"{round(coverage)}%" if coverage is not None else "—", ""])
    trend = report.get("trend_points") or []
    if trend:
        ws.append([])
        ws.append(["Trend zgodności", "", ""])
        ws.cell(row=ws.max_row, column=1).font = HEADER_FONT
        for i, point in enumerate(trend, start=1):
            ws.append([point.get("label") or f"Analiza {i}", round(point.get("score") or 0), ""])
    _autosize(ws)


def render(report: dict, report_type: str = "both") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Podsumowanie"
    _sheet_summary(ws, report)
    _sheet_findings(wb.create_sheet("Ustalenia"), report)
    _sheet_test_cases(wb.create_sheet("Przypadki testowe"), report)
    _sheet_coverage(wb.create_sheet("Pokrycie"), report)

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
